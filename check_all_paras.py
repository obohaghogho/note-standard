import docx

doc = docx.Document("temp_anchor_source/Updated Standard Service Agreement (2).docx")

print("Checking ALL paragraphs for dots, underscores, or blank placeholders...")
for i, p in enumerate(doc.paragraphs):
    text = p.text
    if '…' in text or '...' in text or '____' in text:
        print(f"Para {i:3d}: '{text}'")
